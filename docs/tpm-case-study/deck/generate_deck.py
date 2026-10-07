#!/usr/bin/env python3
"""Generate CareerPilot TPM executive deck (5 slides, 16:9 editable PPTX).

Usage (from repo root, with python-pptx installed):

    .venv/bin/python docs/tpm-case-study/deck/generate_deck.py

Optional PDF (if LibreOffice is available):

    soffice --headless --convert-to pdf --outdir docs/tpm-case-study/deck \\
      docs/tpm-case-study/deck/CareerPilot-TPM-Executive-Deck.pptx
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT_DIR = Path(__file__).resolve().parent
PPTX_PATH = OUT_DIR / "CareerPilot-TPM-Executive-Deck.pptx"

BLACK = RGBColor(0x0A, 0x0A, 0x0A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xF7, 0xF7, 0xFA)
VIOLET = RGBColor(0x7C, 0x3A, 0xED)
VIOLET_SOFT = RGBColor(0xED, 0xE9, 0xFE)
GRAY = RGBColor(0x4B, 0x55, 0x63)
GRAY_LINE = RGBColor(0xE5, 0xE7, 0xEB)
DARK_CARD = RGBColor(0x11, 0x11, 0x18)
MUTED = RGBColor(0x6B, 0x72, 0x80)
GREEN = RGBColor(0x05, 0x96, 0x69)
RED = RGBColor(0xDC, 0x26, 0x26)
RED_BG = RGBColor(0xFE, 0xF2, 0xF2)
BLUE_BG = RGBColor(0xEE, 0xF2, 0xFF)

W, H = Inches(13.333), Inches(7.5)
FOOTER = "CareerPilot — TPM Case Study"


def set_run(run, size=14, bold=False, color=BLACK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def add_round(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    return sh


def textbox(slide, l, t, w, h, text, size=14, bold=False, color=BLACK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def bullets(slide, l, t, w, h, items, size=15, color=BLACK):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = "•  " + item
        set_run(run, size=size, color=color)
    return box


def footer(slide, page, total=5):
    add_rect(slide, Inches(0), Inches(7.15), W, Inches(0.35), BLACK)
    textbox(slide, Inches(0.4), Inches(7.18), Inches(10), Inches(0.3), FOOTER, size=11, color=WHITE)
    textbox(
        slide,
        Inches(11.5),
        Inches(7.18),
        Inches(1.5),
        Inches(0.3),
        f"{page} / {total}",
        size=11,
        color=WHITE,
        align=PP_ALIGN.RIGHT,
    )


def takeaway_bar(slide, text):
    add_round(slide, Inches(0.4), Inches(1.05), Inches(12.5), Inches(0.42), VIOLET_SOFT)
    textbox(
        slide,
        Inches(0.55),
        Inches(1.1),
        Inches(12.2),
        Inches(0.35),
        "Takeaway: " + text,
        size=13,
        bold=True,
        color=VIOLET,
    )


def title_block(slide, title):
    add_rect(slide, Inches(0), Inches(0), W, Inches(0.95), BLACK)
    add_rect(slide, Inches(0), Inches(0.95), Inches(0.12), Inches(6.2), VIOLET)
    textbox(slide, Inches(0.4), Inches(0.28), Inches(12.5), Inches(0.55), title, size=24, bold=True, color=WHITE)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # ---- Slide 1 ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, W, H, WHITE)
    title_block(s, "CareerPilot: Grounded, Human-Controlled Applications")
    takeaway_bar(s, "Ship profile → Fill with evidence honesty — human keeps Submit.")

    add_round(s, Inches(0.4), Inches(1.65), Inches(6.1), Inches(2.0), OFF_WHITE)
    textbox(s, Inches(0.6), Inches(1.75), Inches(5.7), Inches(0.35), "PROBLEM", size=12, bold=True, color=VIOLET)
    bullets(
        s,
        Inches(0.55),
        Inches(2.15),
        Inches(5.8),
        Inches(1.4),
        [
            "Scoring & generation can outrun evidence",
            "Stale packages & cross-user leakage destroy trust",
            "Form automation can cross into Submit",
        ],
        size=14,
    )

    add_round(s, Inches(6.8), Inches(1.65), Inches(6.1), Inches(2.0), OFF_WHITE)
    textbox(s, Inches(7.0), Inches(1.75), Inches(5.7), Inches(0.35), "GOAL", size=12, bold=True, color=VIOLET)
    bullets(
        s,
        Inches(6.95),
        Inches(2.15),
        Inches(5.8),
        Inches(1.4),
        [
            "Grounded workflow to assisted application",
            "Human controls eligibility, EEO/terms, Submit",
            "Local / self-hostable v1 — not hosted SaaS",
        ],
        size=14,
    )

    steps = ["Profile", "Discover", "Fit", "Materials", "Approval", "Assisted Fill", "Track"]
    x = 0.35
    for i, step in enumerate(steps):
        add_round(s, Inches(x), Inches(4.0), Inches(1.55), Inches(0.55), VIOLET if i == 5 else BLACK)
        textbox(
            s,
            Inches(x),
            Inches(4.1),
            Inches(1.55),
            Inches(0.4),
            step,
            size=11,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER,
        )
        if i < len(steps) - 1:
            textbox(
                s,
                Inches(x + 1.48),
                Inches(4.1),
                Inches(0.28),
                Inches(0.4),
                "→",
                size=14,
                bold=True,
                color=VIOLET,
                align=PP_ALIGN.CENTER,
            )
        x += 1.82

    add_round(s, Inches(0.4), Inches(4.85), Inches(12.5), Inches(0.7), RED_BG)
    textbox(
        s,
        Inches(0.6),
        Inches(5.0),
        Inches(12.1),
        Inches(0.45),
        "Never auto-submits  ·  EEO / terms stay manual  ·  Greenhouse + Lever Fill only (certified)",
        size=15,
        bold=True,
        color=RED,
        align=PP_ALIGN.CENTER,
    )
    textbox(
        s,
        Inches(0.4),
        Inches(5.75),
        Inches(12.5),
        Inches(0.9),
        "Program proof: explainable Fit, grounded materials, human approval, assisted Fill — with live acceptance evidence.",
        size=14,
        color=GRAY,
    )
    footer(s, 1)

    # ---- Slide 2 ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, W, H, WHITE)
    title_block(s, "Architecture + Critical Path")
    takeaway_bar(s, "One local app — strict evidence chain, not microservices theater.")

    for i, (label, color) in enumerate(
        [
            ("React / Vite  +  Chrome extension", VIOLET),
            ("FastAPI", BLACK),
            ("SQLite / SQLAlchemy", GRAY),
        ]
    ):
        y = 1.65 + i * 0.7
        add_round(s, Inches(0.4), Inches(y), Inches(5.2), Inches(0.55), color)
        textbox(
            s,
            Inches(0.4),
            Inches(y + 0.1),
            Inches(5.2),
            Inches(0.4),
            label,
            size=14,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER,
        )

    textbox(s, Inches(0.4), Inches(3.9), Inches(5.2), Inches(0.3), "EXTERNALS", size=11, bold=True, color=MUTED)
    for i, label in enumerate(["Job sources", "AI providers", "Greenhouse / Lever"]):
        add_round(s, Inches(0.4 + i * 1.8), Inches(4.25), Inches(1.7), Inches(0.45), OFF_WHITE)
        textbox(
            s,
            Inches(0.4 + i * 1.8),
            Inches(4.32),
            Inches(1.7),
            Inches(0.35),
            label,
            size=11,
            color=BLACK,
            align=PP_ALIGN.CENTER,
        )

    add_round(s, Inches(6.0), Inches(1.65), Inches(6.9), Inches(1.35), OFF_WHITE)
    textbox(s, Inches(6.2), Inches(1.75), Inches(6.5), Inches(0.3), "DATA BOUNDARY", size=11, bold=True, color=VIOLET)
    textbox(
        s,
        Inches(6.2),
        Inches(2.15),
        Inches(6.5),
        Inches(0.7),
        "Shared job catalog  vs  User-scoped private records\n(profile, materials, tracker, analytics, saved searches)",
        size=13,
        color=BLACK,
    )

    add_round(s, Inches(6.0), Inches(3.2), Inches(6.9), Inches(2.0), DARK_CARD)
    textbox(s, Inches(6.2), Inches(3.35), Inches(6.5), Inches(0.3), "CRITICAL PATH", size=11, bold=True, color=VIOLET)
    textbox(
        s,
        Inches(6.2),
        Inches(3.8),
        Inches(6.5),
        Inches(1.2),
        "identity → candidate evidence → posting requirements\n→ deterministic Fit → materials → approval → Fill",
        size=14,
        bold=True,
        color=WHITE,
    )
    textbox(
        s,
        Inches(0.4),
        Inches(5.05),
        Inches(12.5),
        Inches(1.5),
        "Architectural decisions\n• Deterministic Fit does not depend on an LLM\n• Fingerprints invalidate stale evidence / materials\n• ATS identity uses posting IDs — never fuzzy title/company",
        size=14,
        color=BLACK,
    )
    footer(s, 2)

    # ---- Slide 3 ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, W, H, WHITE)
    title_block(s, "Program Plan + Risks")
    takeaway_bar(s, "Small-team workstreams; Fit redesign mid-flight; privacy as a ship gate.")

    ms = [
        "1 Identity\n+ profile",
        "2 Discovery\n+ Fit",
        "3 Materials\n+ approval",
        "4 Assisted\nFill",
        "5 Privacy /\nadversarial",
        "6 Release +\nmaintenance",
    ]
    x = 0.35
    for i, m in enumerate(ms):
        add_round(s, Inches(x), Inches(1.65), Inches(2.0), Inches(0.95), VIOLET if i == 5 else BLACK)
        box = s.shapes.add_textbox(Inches(x + 0.05), Inches(1.78), Inches(1.9), Inches(0.75))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = m
        set_run(run, size=11, bold=True, color=WHITE)
        x += 2.15

    textbox(
        s,
        Inches(0.4),
        Inches(2.85),
        Inches(12),
        Inches(0.3),
        "OWNERSHIP (small team — not a large org)",
        size=11,
        bold=True,
        color=MUTED,
    )
    owners = [
        ("Developer A", "Shell / UI · Prepare · Interview placement"),
        ("Developer B", "Discovery · Verification · ATS · Extension"),
        ("Shared / platform", "Auth · Fit contracts · Privacy · Release gates"),
    ]
    x = 0.4
    for title, desc in owners:
        add_round(s, Inches(x), Inches(3.2), Inches(4.0), Inches(1.0), OFF_WHITE)
        textbox(s, Inches(x + 0.15), Inches(3.3), Inches(3.7), Inches(0.3), title, size=13, bold=True, color=VIOLET)
        textbox(s, Inches(x + 0.15), Inches(3.65), Inches(3.7), Inches(0.45), desc, size=12, color=BLACK)
        x += 4.2

    textbox(s, Inches(0.4), Inches(4.4), Inches(12), Inches(0.3), "TOP RISKS → MITIGATIONS", size=11, bold=True, color=MUTED)
    risks = [
        ("Incomplete posting → misleading Fit", "Verified / Potential + content status"),
        ("Stale evidence reused", "Fingerprints invalidate packages"),
        ("Cross-user privacy leakage", "A9 isolation / IDOR / deletion"),
        ("Extension overreach", "No Submit · EEO/terms manual"),
    ]
    x = 0.4
    for risk, mit in risks:
        add_round(s, Inches(x), Inches(4.75), Inches(3.1), Inches(1.55), WHITE, line=GRAY_LINE)
        textbox(s, Inches(x + 0.12), Inches(4.9), Inches(2.85), Inches(0.7), risk, size=12, bold=True, color=BLACK)
        textbox(s, Inches(x + 0.12), Inches(5.6), Inches(2.85), Inches(0.55), mit, size=12, color=GREEN)
        x += 3.2
    footer(s, 3)

    # ---- Slide 4 ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, W, H, WHITE)
    title_block(s, "Launch Evidence + Release Provenance")
    takeaway_bar(s, "CI is necessary; live ATS/privacy acceptance is decisive — and SHA-bound.")

    add_round(s, Inches(0.4), Inches(1.65), Inches(6.1), Inches(3.35), OFF_WHITE)
    textbox(s, Inches(0.6), Inches(1.8), Inches(5.7), Inches(0.35), "AUTOMATED GATES (tag-time)", size=13, bold=True, color=VIOLET)
    bullets(
        s,
        Inches(0.55),
        Inches(2.25),
        Inches(5.8),
        Inches(2.5),
        [
            "Frontend tests: 226",
            "Extension tests: 109",
            "Mapped paths: 124 valid / 0 missing",
            "CI + dependency audits",
            "Full-history secret scan (Gitleaks)",
        ],
        size=15,
    )

    add_round(s, Inches(6.8), Inches(1.65), Inches(6.1), Inches(3.35), DARK_CARD)
    textbox(s, Inches(7.0), Inches(1.8), Inches(5.7), Inches(0.35), "LIVE / MANUAL ACCEPTANCE", size=13, bold=True, color=VIOLET)
    bullets(
        s,
        Inches(6.95),
        Inches(2.25),
        Inches(5.8),
        Inches(2.5),
        [
            "Greenhouse Fill certified (Chrome 152)",
            "Lever Fill certified — never Submit",
            "A9 isolation / IDOR / account deletion",
            "Phase 6 adversarial: no open P0/P1",
            "Attachment: filename alone ≠ proof",
        ],
        size=15,
        color=WHITE,
    )

    add_round(s, Inches(0.4), Inches(5.2), Inches(12.5), Inches(1.35), BLUE_BG)
    textbox(s, Inches(0.6), Inches(5.35), Inches(12.1), Inches(0.35), "RELEASE PROVENANCE", size=12, bold=True, color=VIOLET)
    textbox(
        s,
        Inches(0.6),
        Inches(5.75),
        Inches(12.1),
        Inches(0.65),
        "Certified runtime 7b6c3ee  ≠  release tag v1.0.0 b73a983\nLater docs/version commits did not inherit runtime certification.",
        size=15,
        bold=True,
        color=BLACK,
    )
    footer(s, 4)

    # ---- Slide 5 ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, W, H, WHITE)
    title_block(s, "Results, Tradeoffs, Next Step")
    takeaway_bar(s, "Honest scope is a quality decision — not a lack of ambition.")

    add_round(s, Inches(0.4), Inches(1.65), Inches(6.1), Inches(2.85), OFF_WHITE)
    textbox(s, Inches(0.6), Inches(1.8), Inches(5.7), Inches(0.35), "SHIPPED", size=13, bold=True, color=GREEN)
    bullets(
        s,
        Inches(0.55),
        Inches(2.25),
        Inches(5.8),
        Inches(2.0),
        [
            "Grounded candidate / job evidence",
            "Deterministic Fit + human approval",
            "Greenhouse / Lever assisted Fill",
            "Private per-user workflows",
            "Tagged immutable v1.0.0",
        ],
        size=14,
    )

    add_round(s, Inches(6.8), Inches(1.65), Inches(6.1), Inches(2.85), RED_BG)
    textbox(s, Inches(7.0), Inches(1.8), Inches(5.7), Inches(0.35), "DELIBERATELY NOT SHIPPED", size=13, bold=True, color=RED)
    bullets(
        s,
        Inches(6.95),
        Inches(2.25),
        Inches(5.8),
        Inches(2.0),
        [
            "Auto-submit / auto-apply",
            "Universal ATS support",
            "Hosted production SaaS",
            "EEO / consent automation",
        ],
        size=14,
    )

    textbox(
        s,
        Inches(0.4),
        Inches(4.7),
        Inches(12.5),
        Inches(0.9),
        "Lessons: live acceptance when CI can’t prove browser behavior · release provenance matters · keeping something out of scope can improve product quality",
        size=14,
        color=GRAY,
    )
    add_round(s, Inches(0.4), Inches(5.55), Inches(12.5), Inches(0.95), BLACK)
    textbox(
        s,
        Inches(0.6),
        Inches(5.75),
        Inches(12.1),
        Inches(0.6),
        "Next: targeted post-v1 correctness + migration maturity — not random feature expansion.",
        size=16,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
    )
    footer(s, 5)

    prs.save(str(PPTX_PATH))
    return PPTX_PATH


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
